"""文档解析：PDF / Word / Excel / TXT"""

import os


def parse_document(file_path):
    """
    根据扩展名自动解析文档，返回纯文本内容。
    支持 .pdf, .docx, .xlsx, .txt
    """
    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    if ext == ".pdf":
        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                pages = [p.extract_text() or "" for p in pdf.pages]
                text = "\n".join(pages)
        except ImportError:
            text = "[提示：请安装 pdfplumber 库以解析 PDF]"

    elif ext == ".docx":
        try:
            from docx import Document

            doc = Document(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            text = "[提示：请安装 python-docx 库以解析 Word]"

    elif ext == ".xlsx":
        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, read_only=True)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) or "" for c in row))
            text = "\n".join(rows)
        except ImportError:
            text = "[提示：请安装 openpyxl 库以解析 Excel]"

    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

    else:
        text = f"不支持的文件格式：{ext}"

    return text
